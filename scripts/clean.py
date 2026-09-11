import csv
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


# This file only does the "cleaning" step.
#
# Input:
#   data/raw/07-2026 Updated LEIE Database.csv
#   data/raw/Copy of Department ... georgia.xlsx
#
# Output:
#   data/cleaned/data_source.csv
#   data/cleaned/import_log.csv
#   data/cleaned/excluded_party.csv
#   data/cleaned/identifier.csv
#   data/cleaned/exclusion_record.csv
#
# Important:
#   This file does not connect to PostgreSQL.
#   It only creates clean CSV files that we can inspect before loading.

ROOT = Path(__file__).resolve().parent.parent
RAW_DATA_DIR = ROOT / "data" / "raw"
CLEANED_DATA_DIR = ROOT / "data" / "cleaned"

LEIE_FILE = RAW_DATA_DIR / "07-2026 Updated LEIE Database.csv"
GEORGIA_FILE = RAW_DATA_DIR / "Copy of Department of Community Health Office Of Inspector General List of Excluded Individuals and Entities as of August 7 2026- georgia.xlsx"


def clean_text(value):
    """Clean text values from CSV/Excel."""
    # This handles missing Excel cells.
    # Example: empty Georgia NPI cell -> ""
    if value is None:
        return ""

    # CSV values are already strings, but Excel values can be numbers or dates.
    # str(value) makes everything safe to clean the same way.
    text = str(value).strip()

    # This handles blank strings.
    # Example: "   " -> ""
    if text == "":
        return ""

    # Some names have repeated spaces.
    # Example: "JESSICA   KRISTINE" -> "JESSICA KRISTINE"
    return " ".join(text.split())


def clean_identifier(value):
    """Clean NPI/UPIN values. Return empty string when the ID is missing."""
    text = clean_text(value)

    # In LEIE, 0000000000 means there is no real NPI.
    # We do not want fake NPI rows in identifier.csv.
    if text == "0000000000":
        return ""

    return text


def parse_date(value):
    """
    Convert source date values to YYYY-MM-DD.

    Source files usually use YYYYMMDD:
      20200319 -> 2020-03-19

    Empty values and 00000000 become empty strings.
    Empty string will become NULL when inserted into PostgreSQL later.
    """
    # Missing source date.
    # Example: empty REINDATE cell -> ""
    if value is None:
        return ""

    # openpyxl may read Excel date cells as datetime objects.
    # Example: datetime(2026, 6, 9, 0, 0) -> "2026-06-09"
    if isinstance(value, datetime):
        return value.date().isoformat()

    # This is here in case Python receives a date object without time.
    if isinstance(value, date):
        return value.isoformat()

    text = str(value).strip()

    # LEIE uses 00000000 for dates that do not exist.
    # Example: REINDATE = 00000000 means not reinstated.
    if text == "" or text == "00000000":
        return ""

    # Sometimes Excel numeric values become strings like "20260807.0".
    # PostgreSQL cannot use that, so we remove the .0 first.
    if text.endswith(".0"):
        text = text[:-2]

    # Normal case:
    #   LEIE EXCLDATE = 20200319
    #   Georgia SANCDATE = 19820415
    try:
        return datetime.strptime(text, "%Y%m%d").date().isoformat()
    except ValueError:
        pass

    # Some Georgia dates are written like YYYYDDMM.
    # Example found in the file:
    #   20223108 probably means 2022-08-31
    try:
        return datetime.strptime(text, "%Y%d%m").date().isoformat()
    except ValueError:
        return ""


def get_party_type(first_name, last_name, business_name):
    """Decide whether this row is a person or a business/entity."""
    # If business_name exists, this row is treated as an entity/business.
    # Example: "#1 MARKETING SERVICE, INC" -> ENTITY
    if business_name:
        return "ENTITY"

    # If first or last name exists, this row is treated as a person.
    # Example: FIRSTNAME = ARTHUR, LASTNAME = DIAMOND -> INDIVIDUAL
    if first_name or last_name:
        return "INDIVIDUAL"

    # This should be rare, but keeps the script from crashing
    # if a row does not clearly look like a person or entity.
    return "UNKNOWN"


def read_cleaned_csv(file_name):
    """Read an existing cleaned CSV. If it does not exist yet, return an empty list."""
    path = CLEANED_DATA_DIR / file_name
    if not path.exists():
        return []

    with path.open(newline="", encoding="utf-8") as file:
        return list(csv.DictReader(file))


def max_existing_id(file_name, id_column):
    """Find the largest ID already used in a cleaned CSV."""
    rows = read_cleaned_csv(file_name)
    if not rows:
        return 0

    return max(int(row[id_column]) for row in rows if row[id_column])


def append_csv(file_name, fieldnames, rows):
    """Append rows to one cleaned CSV without deleting old rows."""
    if not rows:
        return

    # Create data/cleaned if it does not already exist.
    CLEANED_DATA_DIR.mkdir(parents=True, exist_ok=True)
    output_path = CLEANED_DATA_DIR / file_name

    # If the file is new or empty, write the header first.
    # If the file already exists, only append new data rows.
    should_write_header = not output_path.exists() or output_path.stat().st_size == 0

    # "a" means append mode.
    # This is the important change for monthly updates:
    # old cleaned rows stay in the file, new rows are added to the bottom.
    with output_path.open("a", newline="", encoding="utf-8") as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames)

        if should_write_header:
            writer.writeheader()

        writer.writerows(rows)


def clean_leie(data_source_id, import_log_id, start_party_id, start_identifier_id, start_exclusion_record_id):
    """Clean LEIE CSV rows and return rows for our database tables."""
    # data_source_id and import_log_id are passed in from main().
    # That allows each monthly source file to get a new ID.

    # Each list below becomes one cleaned CSV file later.
    parties = []
    identifiers = []
    exclusions = []

    # These counters create IDs for the cleaned CSV files.
    # Example: first LEIE row gets party_id = 1.
    party_id = start_party_id
    identifier_id = start_identifier_id
    exclusion_record_id = start_exclusion_record_id

    # utf-8-sig handles CSV files that may start with a hidden BOM marker.
    with LEIE_FILE.open(newline="", encoding="utf-8-sig") as file:
        # DictReader lets us read by column name.
        # Example: row["FIRSTNAME"], row["NPI"], row["EXCLDATE"]
        reader = csv.DictReader(file)

        for row in reader:
            # LEIE separates person names and business names.
            # If LASTNAME/FIRSTNAME are empty and BUSNAME exists, it is an entity.
            first_name = clean_text(row["FIRSTNAME"])
            middle_name = clean_text(row["MIDNAME"])
            last_name = clean_text(row["LASTNAME"])
            business_name = clean_text(row["BUSNAME"])

            # excluded_party table: who is excluded?
            #
            # Field mapping:
            #   LASTNAME  -> last_name
            #   FIRSTNAME -> first_name
            #   MIDNAME   -> middle_name
            #   BUSNAME   -> business_name
            #   GENERAL   -> provider_category
            #   SPECIALTY -> specialty
            #   DOB       -> dob
            #   ADDRESS/CITY/STATE/ZIP -> address fields
            parties.append({
                "party_id": party_id,
                "party_type": get_party_type(first_name, last_name, business_name),
                "first_name": first_name,
                "middle_name": middle_name,
                "last_name": last_name,
                "business_name": business_name,
                "provider_category": clean_text(row["GENERAL"]),
                "specialty": clean_text(row["SPECIALTY"]),
                "dob": parse_date(row["DOB"]),
                "address": clean_text(row["ADDRESS"]),
                "city": clean_text(row["CITY"]),
                "state": clean_text(row["STATE"]),
                "zip_code": clean_text(row["ZIP"]),
            })

            # identifier table: NPI and UPIN, only when real values exist.
            #
            # Many LEIE rows have NPI = 0000000000.
            # That means no NPI, so those rows are not added to identifier.csv.
            npi = clean_identifier(row["NPI"])
            if npi:
                identifiers.append({
                    "identifier_id": identifier_id,
                    "party_id": party_id,
                    "identifier_type": "NPI",
                    "identifier_value": npi,
                })
                identifier_id += 1

            # UPIN is an older provider identifier.
            # Most rows may not have it, but we keep it when present.
            upin = clean_identifier(row["UPIN"])
            if upin:
                identifiers.append({
                    "identifier_id": identifier_id,
                    "party_id": party_id,
                    "identifier_type": "UPIN",
                    "identifier_value": upin,
                })
                identifier_id += 1

            # exclusion_record table: what exclusion happened?
            #
            # Field mapping:
            #   EXCLTYPE   -> exclusion_type
            #   EXCLDATE   -> exclusion_date
            #   REINDATE   -> reinstatement_date
            #   WAIVERDATE -> waiver_date
            #   WVRSTATE   -> waiver_state
            #
            # REINDATE and WAIVERDATE are often 00000000,
            # so parse_date() turns them into "".
            exclusions.append({
                "exclusion_record_id": exclusion_record_id,
                "party_id": party_id,
                "data_source_id": data_source_id,
                "import_log_id": import_log_id,
                "exclusion_type": clean_text(row["EXCLTYPE"]),
                "exclusion_date": parse_date(row["EXCLDATE"]),
                "reinstatement_date": parse_date(row["REINDATE"]),
                "waiver_date": parse_date(row["WAIVERDATE"]),
                "waiver_state": clean_text(row["WVRSTATE"]),
                "status": "ACTIVE",
            })

            # Move to the next generated IDs for the next source row.
            party_id += 1
            exclusion_record_id += 1

    return parties, identifiers, exclusions, party_id, identifier_id, exclusion_record_id


def clean_georgia(data_source_id, import_log_id, start_party_id, start_identifier_id, start_exclusion_record_id):
    """Clean Georgia Excel rows and return rows for our database tables."""
    # data_source_id and import_log_id are passed in from main().
    # That allows each monthly source file to get a new ID.

    # Each list below becomes part of a cleaned CSV file.
    parties = []
    identifiers = []
    exclusions = []

    # Start from the next IDs after the LEIE rows.
    # This avoids duplicate party_id / identifier_id / exclusion_record_id.
    party_id = start_party_id
    identifier_id = start_identifier_id
    exclusion_record_id = start_exclusion_record_id

    # Load the Excel workbook.
    # read_only=True makes reading faster.
    # data_only=True reads the displayed cell value.
    workbook = load_workbook(GEORGIA_FILE, read_only=True, data_only=True)
    worksheet = workbook["Sheet1"]

    # Georgia file layout:
    # row 1 = title
    # row 2 = blank
    # row 3 = real column names
    headers = list(next(worksheet.iter_rows(min_row=3, max_row=3, values_only=True))[:8])

    # Data starts at row 4 because row 3 is the header.
    for values in worksheet.iter_rows(min_row=4, values_only=True):
        # Convert the Excel row into a dictionary.
        # Example: row["LAST NAME"], row["GENERAL"], row["SANCDATE"]
        row = dict(zip(headers, values[:8]))
        exclusion_date = parse_date(row["SANCDATE"])

        # Our exclusion_record table requires exclusion_date.
        # If Georgia has no SANCDATE, skip that row.
        #
        # This handles rows in the Georgia file where SANCDATE is blank.
        if not exclusion_date:
            continue

        # Georgia uses different column names from LEIE.
        # Example:
        #   "FIRST NAME" instead of "FIRSTNAME"
        #   "MIDDLE NAME" instead of "MIDNAME"
        first_name = clean_text(row["FIRST NAME"])
        middle_name = clean_text(row["MIDDLE NAME"])
        last_name = clean_text(row["LAST NAME"])
        business_name = clean_text(row["BUSINESS NAME"])

        # Georgia does not include DOB/address/city/zip in this file,
        # so those fields become "" in the cleaned CSV.
        #
        # Georgia GENERAL still describes the provider type/category,
        # so it maps to provider_category.
        parties.append({
            "party_id": party_id,
            "party_type": get_party_type(first_name, last_name, business_name),
            "first_name": first_name,
            "middle_name": middle_name,
            "last_name": last_name,
            "business_name": business_name,
            "provider_category": clean_text(row["GENERAL"]),
            "specialty": "",
            "dob": "",
            "address": "",
            "city": "",
            "state": clean_text(row["STATE"]),
            "zip_code": "",
        })

        # Georgia only has NPI, not UPIN.
        # If NPI is blank, no identifier row is created.
        npi = clean_identifier(row["NPI"])
        if npi:
            identifiers.append({
                "identifier_id": identifier_id,
                "party_id": party_id,
                "identifier_type": "NPI",
                "identifier_value": npi,
            })
            identifier_id += 1

        # Georgia does not have EXCLTYPE, REINDATE, WAIVERDATE, or WVRSTATE.
        # Those fields become "" in the cleaned exclusion_record.csv.
        #
        # Georgia SANCDATE maps to exclusion_date.
        exclusions.append({
            "exclusion_record_id": exclusion_record_id,
            "party_id": party_id,
            "data_source_id": data_source_id,
            "import_log_id": import_log_id,
            "exclusion_type": "",
            "exclusion_date": exclusion_date,
            "reinstatement_date": "",
            "waiver_date": "",
            "waiver_state": "",
            "status": "ACTIVE",
        })

        # Move to the next generated IDs for the next Georgia row.
        party_id += 1
        exclusion_record_id += 1

    return parties, identifiers, exclusions


def main():
    # Read existing cleaned files first.
    # This is how the script knows where to continue numbering IDs.
    existing_data_sources = read_cleaned_csv("data_source.csv")
    processed_file_names = {row["file_name"] for row in existing_data_sources}

    next_data_source_id = max_existing_id("data_source.csv", "data_source_id") + 1
    next_import_log_id = max_existing_id("import_log.csv", "import_log_id") + 1
    next_party_id = max_existing_id("excluded_party.csv", "party_id") + 1
    next_identifier_id = max_existing_id("identifier.csv", "identifier_id") + 1
    next_exclusion_id = max_existing_id("exclusion_record.csv", "exclusion_record_id") + 1

    # These rows become new rows in data_source.csv.
    #
    # data_source.csv answers:
    #   Which official file did this data come from?
    #
    data_sources = []
    import_logs = []
    all_parties = []
    all_identifiers = []
    all_exclusions = []

    # Clean LEIE only if this exact file name has not been processed before.
    # This prevents duplicate rows when you run clean.py multiple times.
    if LEIE_FILE.name in processed_file_names:
        print(f"Skipped already-cleaned file: {LEIE_FILE.name}")
    else:
        leie_data_source_id = next_data_source_id
        leie_import_log_id = next_import_log_id

        leie_parties, leie_identifiers, leie_exclusions, next_party_id, next_identifier_id, next_exclusion_id = clean_leie(
            data_source_id=leie_data_source_id,
            import_log_id=leie_import_log_id,
            start_party_id=next_party_id,
            start_identifier_id=next_identifier_id,
            start_exclusion_record_id=next_exclusion_id,
        )

        data_sources.append({
            "data_source_id": leie_data_source_id,
            "source_name": "HHS OIG LEIE",
            "file_name": LEIE_FILE.name,
            "file_type": "CSV",
            "source_date": "2026-07-01",
        })
        import_logs.append({
            "import_log_id": leie_import_log_id,
            "data_source_id": leie_data_source_id,
            "imported_at": datetime.now().isoformat(timespec="seconds"),
            "status": "SUCCESS",
            "records_loaded": len(leie_exclusions),
            "notes": "Cleaned from LEIE CSV",
        })
        all_parties += leie_parties
        all_identifiers += leie_identifiers
        all_exclusions += leie_exclusions

        next_data_source_id += 1
        next_import_log_id += 1

    # Clean Georgia only if this exact file name has not been processed before.
    if GEORGIA_FILE.name in processed_file_names:
        print(f"Skipped already-cleaned file: {GEORGIA_FILE.name}")
    else:
        georgia_data_source_id = next_data_source_id
        georgia_import_log_id = next_import_log_id

        georgia_parties, georgia_identifiers, georgia_exclusions = clean_georgia(
            data_source_id=georgia_data_source_id,
            import_log_id=georgia_import_log_id,
            start_party_id=next_party_id,
            start_identifier_id=next_identifier_id,
            start_exclusion_record_id=next_exclusion_id,
        )

        data_sources.append({
            "data_source_id": georgia_data_source_id,
            "source_name": "Georgia DCH OIG",
            "file_name": GEORGIA_FILE.name,
            "file_type": "XLSX",
            "source_date": "2026-08-07",
        })
        import_logs.append({
            "import_log_id": georgia_import_log_id,
            "data_source_id": georgia_data_source_id,
            "imported_at": datetime.now().isoformat(timespec="seconds"),
            "status": "SUCCESS",
            "records_loaded": len(georgia_exclusions),
            "notes": "Cleaned from Georgia XLSX; rows without SANCDATE skipped",
        })
        all_parties += georgia_parties
        all_identifiers += georgia_identifiers
        all_exclusions += georgia_exclusions

    # Append new rows to the five cleaned CSV files.
    # These file names match the five PostgreSQL tables.
    append_csv(
        "data_source.csv",
        ["data_source_id", "source_name", "file_name", "file_type", "source_date"],
        data_sources,
    )
    append_csv(
        "import_log.csv",
        ["import_log_id", "data_source_id", "imported_at", "status", "records_loaded", "notes"],
        import_logs,
    )
    append_csv(
        "excluded_party.csv",
        [
            "party_id", "party_type", "first_name", "middle_name", "last_name",
            "business_name", "provider_category", "specialty", "dob", "address",
            "city", "state", "zip_code",
        ],
        all_parties,
    )
    append_csv(
        "identifier.csv",
        ["identifier_id", "party_id", "identifier_type", "identifier_value"],
        all_identifiers,
    )
    append_csv(
        "exclusion_record.csv",
        [
            "exclusion_record_id", "party_id", "data_source_id", "import_log_id",
            "exclusion_type", "exclusion_date", "reinstatement_date", "waiver_date",
            "waiver_state", "status",
        ],
        all_exclusions,
    )

    print("New cleaned rows appended to data/cleaned")
    print(f"new data_source rows: {len(data_sources)}")
    print(f"new import_log rows: {len(import_logs)}")
    print(f"new excluded_party rows: {len(all_parties)}")
    print(f"new identifier rows: {len(all_identifiers)}")
    print(f"new exclusion_record rows: {len(all_exclusions)}")


if __name__ == "__main__":
    main()
